"""Rasterize the poster .pptx to PNG (1920x1080) without PowerPoint.

Supports exactly the shape vocabulary build_poster.py emits: slide background,
text boxes (runs with size/bold/italic/color, paragraph alignment, vertical
anchor), filled rectangles, one native table, and the movie (drawn as its poster
frame, or left as a hole when --hole is given so ffmpeg can overlay the video).
Also writes layout.json with the movie rectangle in pixels for the overlay step.
"""
import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu

FONT_TTC = '/System/Library/Fonts/HelveticaNeue.ttc'
# face indices inside HelveticaNeue.ttc (0 regular, 1 bold, 2 italic ... verified at runtime)


def face_index(bold, italic):
    want = ('Bold' if bold else 'Regular') + (' Italic' if italic else '')
    for i in range(0, 14):
        try:
            f = ImageFont.truetype(FONT_TTC, 20, index=i)
            fam, sty = f.getname()
            if sty.replace('Oblique', 'Italic') == want:
                return i
        except Exception:
            break
    return 1 if bold else 0


FACE = {(b, i): face_index(b, i) for b in (False, True) for i in (False, True)}


def font(pt, bold=False, italic=False, scale=1.0):
    return ImageFont.truetype(FONT_TTC, max(1, int(round(pt * scale))), index=FACE[(bold, italic)])


def rgb_of(color):
    try:
        if color.type is not None:
            return tuple(int(x) for x in bytes.fromhex(str(color.rgb)))
    except Exception:
        pass
    return (20, 20, 20)


def wrap(draw, text, fnt, max_w):
    words = text.split(' ')
    lines, cur = [], ''
    for w in words:
        t = (cur + ' ' + w).strip()
        if draw.textlength(t, font=fnt) <= max_w or not cur:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_text_frame(draw, tf, box, px_per_in):
    x, y, w, h = box
    ml = Emu(tf.margin_left).inches * px_per_in; mr = Emu(tf.margin_right).inches * px_per_in
    mt = Emu(tf.margin_top).inches * px_per_in; mb = Emu(tf.margin_bottom).inches * px_per_in
    inner_w = w - ml - mr
    pt_scale = px_per_in / 72.0
    blocks = []
    for p in tf.paragraphs:
        if not p.runs:
            blocks.append(([], None, p.alignment, 0)); continue
        r = p.runs[0]
        size = r.font.size.pt if r.font.size else 18
        fnt = font(size, bool(r.font.bold), bool(r.font.italic), pt_scale)
        text = ''.join(rr.text for rr in p.runs)
        lines = wrap(draw, text, fnt, inner_w)
        lh = fnt.size * 1.18
        blocks.append((lines, fnt, p.alignment, lh, rgb_of(r.font.color)))
    total_h = sum(len(b[0]) * b[3] if b[0] else b[3] for b in blocks)
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = y + mt + (h - mt - mb - total_h) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = y + h - mb - total_h
    else:
        cy = y + mt
    for b in blocks:
        if not b[0]:
            cy += b[3]; continue
        lines, fnt, align, lh, col = b
        for line in lines:
            tw = draw.textlength(line, font=fnt)
            if align == PP_ALIGN.CENTER:
                lx = x + ml + (inner_w - tw) / 2
            elif align == PP_ALIGN.RIGHT:
                lx = x + w - mr - tw
            else:
                lx = x + ml
            draw.text((lx, cy), line, font=fnt, fill=col)
            cy += lh


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('pptx'); ap.add_argument('out_png')
    ap.add_argument('--width', type=int, default=1920)
    ap.add_argument('--hole', action='store_true', help='leave the movie area as paper (for video overlay)')
    args = ap.parse_args()
    prs = Presentation(args.pptx)
    s = prs.slides[0]
    W = args.width; px_per_in = W / Emu(prs.slide_width).inches
    H = int(round(Emu(prs.slide_height).inches * px_per_in))
    bg = (246, 245, 241)
    try:
        bg = rgb_of(s.background.fill.fore_color)
    except Exception:
        pass
    img = Image.new('RGB', (W, H), bg)
    draw = ImageDraw.Draw(img)
    layout = {'width': W, 'height': H, 'movie': None}

    def box(sh):
        return (Emu(sh.left).inches * px_per_in, Emu(sh.top).inches * px_per_in,
                Emu(sh.width).inches * px_per_in, Emu(sh.height).inches * px_per_in)

    for sh in s.shapes:
        x, y, w, h = box(sh)
        if sh.shape_type == MSO_SHAPE_TYPE.MEDIA or 'replay' in (sh.name or '').lower():
            layout['movie'] = {'x': int(round(x)), 'y': int(round(y)), 'w': int(round(w)), 'h': int(round(h))}
            if not args.hole:
                try:
                    blip = sh._element.xpath('.//a:blip/@r:embed')[0]
                    part = sh.part.related_part(blip)
                    from io import BytesIO
                    poster = Image.open(BytesIO(part.blob)).convert('RGB').resize((int(w), int(h)))
                    img.paste(poster, (int(x), int(y)))
                except Exception as e:
                    draw.rectangle([x, y, x + w, y + h], outline=(200, 200, 200))
            continue
        if sh.has_table if hasattr(sh, 'has_table') else False:
            tbl = sh.table
            cy = y
            for r in range(len(tbl.rows)):
                cx = x
                rh = Emu(tbl.rows[r].height).inches * px_per_in
                for c in range(len(tbl.columns)):
                    cw = Emu(tbl.columns[c].width).inches * px_per_in
                    cell = tbl.cell(r, c)
                    fill = None
                    try:
                        if cell.fill.type is not None:
                            fill = rgb_of(cell.fill.fore_color)
                    except Exception:
                        pass
                    if fill:
                        draw.rectangle([cx, cy, cx + cw, cy + rh], fill=fill)
                    draw.rectangle([cx, cy, cx + cw, cy + rh], outline=(246, 245, 241), width=2)
                    draw_text_frame(draw, cell.text_frame, (cx, cy, cw, rh), px_per_in)
                    cx += cw
                cy += rh
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                if sh.fill.type is not None:
                    draw.rectangle([x, y, x + w, y + h], fill=rgb_of(sh.fill.fore_color))
            except Exception:
                pass
            try:
                from pptx.enum.dml import MSO_FILL
                if sh.line.fill.type == MSO_FILL.SOLID:
                    draw.rectangle([x, y, x + w, y + h], outline=rgb_of(sh.line.color), width=1)
            except Exception:
                pass
        if sh.has_text_frame and sh.text_frame.text.strip():
            draw_text_frame(draw, sh.text_frame, (x, y, w, h), px_per_in)
    img.save(args.out_png)
    json.dump(layout, open(Path(args.out_png).with_suffix('.layout.json'), 'w'))
    print('wrote', args.out_png, layout['movie'])


if __name__ == '__main__':
    main()
